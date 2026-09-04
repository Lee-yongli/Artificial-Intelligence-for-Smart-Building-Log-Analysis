from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE_DIR = Path("/Users/yonglijiang/Documents/Adelaide University/AU Course/AU Course - 2026 Semester 2/Industry Research Project/2 Proposal Report")
SOURCE = BASE_DIR / "individual_8min_speech_yongli_jiang.md"
OUTPUT = BASE_DIR / "individual_8min_speech_yongli_jiang.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(85, 85, 85)
BLACK = RGBColor(0, 0, 0)


def set_run_font(run, name="Calibri", east_asia="PingFang SC", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def paragraph_bottom_border(paragraph, color="D9E2F3", size="8", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_title_block(doc):
    title = doc.add_paragraph()
    set_paragraph_spacing(title, before=0, after=4, line=1.1)
    run = title.add_run("Individual Speech Script")
    set_run_font(run, size=22, color=BLACK, bold=True)

    subtitle = doc.add_paragraph()
    set_paragraph_spacing(subtitle, before=0, after=12, line=1.15)
    run = subtitle.add_run("AI Smart-Building Log Analysis")
    set_run_font(run, size=13, color=MUTED)

    meta = doc.add_paragraph()
    set_paragraph_spacing(meta, before=0, after=12, line=1.15)
    run = meta.add_run("Speaker: Yongli Jiang (Lee) | Role: Data Engineer | Target: 5 to 7 minutes with pauses")
    set_run_font(run, size=10.5, color=MUTED)
    paragraph_bottom_border(meta)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading1 = doc.styles["Heading 1"]
    heading1.font.name = "Calibri"
    heading1.font.size = Pt(16)
    heading1.font.color.rgb = BLUE
    heading1.font.bold = True
    heading1._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    heading1._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    heading1._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    heading1.paragraph_format.space_before = Pt(18)
    heading1.paragraph_format.space_after = Pt(10)
    heading1.paragraph_format.line_spacing = 1.25

    heading2 = doc.styles["Heading 2"]
    heading2.font.name = "Calibri"
    heading2.font.size = Pt(13)
    heading2.font.color.rgb = BLUE
    heading2.font.bold = True
    heading2._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    heading2._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    heading2._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    heading2.paragraph_format.space_before = Pt(14)
    heading2.paragraph_format.space_after = Pt(7)
    heading2.paragraph_format.line_spacing = 1.25


def clean_line(line):
    return line.rstrip().replace("  ", " ").strip()


def add_qna_question(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=6, after=3, line=1.25)
    run = p.add_run(text)
    set_run_font(run, size=11, color=DARK_BLUE, bold=True)


def add_body_paragraph(doc, text, italic=False):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=6, line=1.25)
    run = p.add_run(text)
    set_run_font(run, size=11, color=BLACK, italic=italic)


def build_docx():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_title_block(doc)

    skip_meta = True
    in_qna = False
    pending = []

    def flush_pending():
        if not pending:
            return
        text = " ".join(pending).strip()
        pending.clear()
        if text:
            add_body_paragraph(doc, text)

    for raw in lines:
        line = clean_line(raw)
        if not line:
            flush_pending()
            continue

        if line.startswith("# Individual "):
            continue
        if re.match(r"^(Target speaker|Role|Target duration|Style):", line):
            continue

        if line.startswith("## "):
            flush_pending()
            heading = line[3:].strip()
            if heading == "Short Q&A Preparation":
                doc.add_page_break()
                in_qna = True
            doc.add_paragraph(heading, style="Heading 1")
            continue

        if in_qna and line.startswith('If asked "'):
            flush_pending()
            add_qna_question(doc, line)
            continue

        if in_qna and line.startswith("Chinese reminder:"):
            flush_pending()
            add_body_paragraph(
                doc,
                "Reminder: first say it is not a traditional text-based event log. "
                "Then explain that it is still time-stamped operational data. "
                "Finally, mention that future work should validate the prototype with real logs or maintenance records.",
                italic=True,
            )
            continue

        pending.append(line)

    flush_pending()

    doc.core_properties.title = "Individual Speech Script - Yongli Jiang"
    doc.core_properties.subject = "AI Smart-Building Log Analysis"
    doc.core_properties.author = "Yongli Jiang"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
